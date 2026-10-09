import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="MedShield_Drug_Risk_Analyzer_Presentation.pptx"):
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    C_NAVY = RGBColor(15, 23, 42)        # #0F172A Dark Slate
    C_SLATE = RGBColor(30, 41, 59)       # #1E293B
    C_BLUE = RGBColor(37, 99, 235)       # #2563EB Primary Blue
    C_CYAN = RGBColor(14, 165, 233)      # #0EA5E9 Accent
    C_RED = RGBColor(239, 68, 68)        # #EF4444 Critical Alert
    C_GREEN = RGBColor(16, 185, 129)     # #10B981 Suppressed Safe
    C_AMBER = RGBColor(245, 158, 11)     # #F59E0B Warning
    C_BG_LIGHT = RGBColor(248, 250, 252) # #F8FAFC Card BG
    C_BORDER = RGBColor(226, 232, 240)   # #E2E8F0
    C_MUTED = RGBColor(100, 116, 139)    # #64748B
    C_DARK_TEXT = RGBColor(30, 41, 59)   # #1E293B
    C_WHITE = RGBColor(255, 255, 255)

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(245, 247, 250)
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text):
        # Category pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.8), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(224, 231, 255)
        pill.line.fill.background()
        p_tf = pill.text_frame
        p_tf.word_wrap = True
        p_para = p_tf.paragraphs[0]
        p_para.text = category_text.upper()
        p_para.font.size = Pt(9.5)
        p_para.font.bold = True
        p_para.font.color.rgb = C_BLUE
        p_para.alignment = PP_ALIGN.CENTER

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        para = tf.paragraphs[0]
        para.text = title_text
        para.font.size = Pt(22)
        para.font.bold = True
        para.font.color.rgb = C_NAVY

    # =========================================================================
    # SLIDE 1: EXECUTIVE SUMMARY & THE ALERT FATIGUE CHALLENGE
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_bg(slide1)
    add_header(slide1, "Eliminating EHR Alert Fatigue with Contextual Clinical AI", "Executive Overview & Problem Statement")

    # Hero card on Left (The Clinical Dilemma)
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.4))
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_WHITE
    card1.line.color.rgb = C_BORDER
    tf1 = card1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.35)
    tf1.margin_right = Inches(0.35)
    tf1.margin_top = Inches(0.35)

    p = tf1.paragraphs[0]
    p.text = "THE CRISIS: CLINICIAN ALERT FATIGUE"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_RED

    points_c1 = [
        ("90%+ Alert Override Rate: ", "Traditional EHR systems blindly trigger pop-ups for any drug pair, regardless of patient reality. Clinicians become desensitized and bypass warnings."),
        ("Missed Life-Threatening Risks: ", "When severe interactions occur in vulnerable patients (e.g. renal failure, abnormal INR), warnings get lost in the noise, leading to preventable hospitalizations."),
        ("Zero Patient Lab Context: ", "Legacy rule-engines do not cross-reference organ clearance (eGFR), metabolic enzymes (ALT/AST), or current electrolyte balances (Potassium, INR).")
    ]
    for bold_prefix, text in points_c1:
        p2 = tf1.add_paragraph()
        p2.space_before = Pt(14)
        run1 = p2.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(11)
        run1.font.color.rgb = C_NAVY
        run2 = p2.add_run()
        run2.text = text
        run2.font.size = Pt(10.5)
        run2.font.color.rgb = C_SLATE

    # Stat highlight at bottom of card 1
    stat_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.15), Inches(5.4), Inches(4.9), Inches(1.1))
    stat_box.fill.solid()
    stat_box.fill.fore_color.rgb = RGBColor(254, 242, 242)
    stat_box.line.color.rgb = RGBColor(254, 202, 202)
    stf = stat_box.text_frame
    stf.word_wrap = True
    stf.margin_top = Inches(0.18)
    sp = stf.paragraphs[0]
    sp.text = "🚨 Over 90% of Hospital DDI Alerts Ignored"
    sp.font.bold = True
    sp.font.size = Pt(13)
    sp.font.color.rgb = C_RED
    sp.alignment = PP_ALIGN.CENTER
    sp2 = stf.add_paragraph()
    sp2.text = "Direct cause of medication errors & physician burnout"
    sp2.font.size = Pt(10)
    sp2.font.color.rgb = C_SLATE
    sp2.alignment = PP_ALIGN.CENTER

    # Right Column: The Solution (MedShield AI)
    card2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.4))
    card2.fill.solid()
    card2.fill.fore_color.rgb = C_WHITE
    card2.line.color.rgb = C_BORDER
    tf2 = card2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0.35)
    tf2.margin_right = Inches(0.35)
    tf2.margin_top = Inches(0.35)

    p = tf2.paragraphs[0]
    p.text = "THE SOLUTION: PRECISION CONTEXTUAL CDS"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    pillars = [
        ("1. FHIR Patient Laboratory RAG: ", "Ingests patient-specific eGFR, serum creatinine, potassium, INR, and transaminases directly into the reasoning prompt."),
        ("2. Intelligent Alert Fatigue Suppression: ", "Actively suppresses generic alerts when patient organ function is normal and therapy is benign (e.g. statin + short-course azithromycin)."),
        ("3. Actionable Clinical Decision Bar: ", "Empowers clinicians with 1-click evidence-based substitutions (e.g. Apixaban), dose adjustments, or reasoned overrides."),
        ("4. Immutable HIPAA Audit Trail: ", "Tamper-evident logging of every clinical action, rationale, and provider ID for regulatory governance.")
    ]
    for bold_prefix, text in pillars:
        p2 = tf2.add_paragraph()
        p2.space_before = Pt(12)
        run1 = p2.add_run()
        run1.text = bold_prefix
        run1.font.bold = True
        run1.font.size = Pt(11)
        run1.font.color.rgb = C_NAVY
        run2 = p2.add_run()
        run2.text = text
        run2.font.size = Pt(10.5)
        run2.font.color.rgb = C_SLATE

    # =========================================================================
    # SLIDE 2: SYSTEM ARCHITECTURE & CONTEXTUAL RAG FLOW
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_bg(slide2)
    add_header(slide2, "End-to-End Contextual RAG Architecture & Decision Engine", "Architecture & Data Flow")

    # Helper function for flow nodes
    def make_node(slide, left, top, width, height, title, subtitle, bg_color, border_color, title_color=C_NAVY):
        node = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        node.fill.solid()
        node.fill.fore_color.rgb = bg_color
        node.line.color.rgb = border_color
        node.line.width = Pt(1.5)
        tf = node.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.1)
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = title_color
        p.alignment = PP_ALIGN.CENTER
        
        if subtitle:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(2)
            p2.text = subtitle
            p2.font.size = Pt(8.5)
            p2.font.color.rgb = C_MUTED
            p2.alignment = PP_ALIGN.CENTER
        return node

    # Top Node: Clinician Prescription
    make_node(slide2, 4.6, 1.4, 4.1, 0.75, "Clinician Prescription Order", "Patient Profile + Newly Ordered Medication", C_WHITE, C_BLUE, C_BLUE)

    # Middle Layer: FHIR Record, RAG Retriever, ChromaDB
    make_node(slide2, 0.8, 2.5, 3.4, 0.95, "FHIR Patient Record", "Labs (eGFR, INR, K+), Diagnoses,\nActive Medications & Dosages", RGBColor(240, 249, 255), RGBColor(186, 230, 253), C_NAVY)
    make_node(slide2, 4.8, 2.5, 3.7, 0.95, "RAG Engine Retriever", "Semantic Vector Search &\nPairwise Interaction Matcher", C_WHITE, C_BLUE, C_BLUE)
    make_node(slide2, 9.1, 2.5, 3.4, 0.95, "Local ChromaDB Store", "DrugBank / RxNorm Chunks\nMechanisms, Enzymes, Overrides", RGBColor(245, 243, 255), RGBColor(221, 214, 254), RGBColor(109, 40, 217))

    # Core LLM Evaluator
    make_node(slide2, 4.3, 3.8, 4.7, 0.85, "LLM Contextual Evaluator (MaaS)", "azure/genailab-maas-gpt-4o-mini | Temperature 0.1\nStrict JSON Schema & Suppression Rules", RGBColor(254, 243, 199), RGBColor(252, 211, 77), RGBColor(180, 83, 9))

    # Bifurcation: Contextual Risk Active vs Contextually Suppressed
    make_node(slide2, 1.4, 4.95, 4.8, 0.95, "🔴 Contextual Risk Active (Tier 1/2)", "Elevates to Critical/Moderate\nIdentifies Abnormal Lab Driver (INR 3.2, eGFR 34)", RGBColor(254, 242, 242), C_RED, C_RED)
    make_node(slide2, 7.1, 4.95, 4.8, 0.95, "🟢 Contextually Suppressed (Tier 3)", "Alert Fatigue Eliminated (is_suppressed=true)\nNormal Organ Reserve & Lab Markers", RGBColor(240, 253, 244), C_GREEN, C_GREEN)

    # Clinician Action & Decision
    make_node(slide2, 3.8, 6.1, 5.7, 0.55, "Clinician Action & Decision Workflow", "Accept Alternate Drug  |  Adjust Dosage  |  Override with Clinical Reason", C_WHITE, C_SLATE, C_NAVY)

    # Base: HIPAA Audit Trail Log
    make_node(slide2, 4.3, 6.8, 4.7, 0.45, "Immutable HIPAA Audit Trail (audit_log.json)", "Records Timestamp, Patient ID, Drug, Action & Provider ID", RGBColor(241, 245, 249), C_BORDER, C_SLATE)

    # Arrows / Connecting indicator callout labels
    callout = slide2.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(2.7), Inches(0.9))
    ctf = callout.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = "⚡ Key Innovation:"
    cp.font.bold = True
    cp.font.size = Pt(9.5)
    cp.font.color.rgb = C_BLUE
    cp2 = ctf.add_paragraph()
    cp2.text = "Patient lab reality dictates whether an alert is shown or suppressed."
    cp2.font.size = Pt(8.5)
    cp2.font.color.rgb = C_SLATE

    # =========================================================================
    # SLIDE 3: CLINICAL HIGHLIGHTS, FEATURES & VALIDATION
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_bg(slide3)
    add_header(slide3, "Key Clinical Highlights, Real-World Cases & Governance", "Features, Validation & Governance")

    # 3 Columns for features / cases
    col_w = Inches(3.7)
    gap = Inches(0.3)
    c1_left = Inches(0.8)
    c2_left = c1_left + col_w + gap
    c3_left = c2_left + col_w + gap

    # Card A: Severe Contextual Risk
    ca = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, Inches(1.5), col_w, Inches(5.4))
    ca.fill.solid()
    ca.fill.fore_color.rgb = C_WHITE
    ca.line.color.rgb = RGBColor(254, 202, 202)
    ca_tf = ca.text_frame
    ca_tf.word_wrap = True
    ca_tf.margin_top = Inches(0.3)
    ca_tf.margin_left = Inches(0.25)
    ca_tf.margin_right = Inches(0.25)

    p = ca_tf.paragraphs[0]
    p.text = "🔴 CASE 1: SEVERE RISK"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_RED

    points_a = [
        ("Patient: ", "Eleanor Vance (72F), Atrial Fibrillation, CKD Stage 3b"),
        ("Active Regimen: ", "Warfarin 5mg daily"),
        ("Prescribed: ", "Amiodarone 200mg daily"),
        ("Lab Context: ", "INR 3.2 (Supratherapeutic), eGFR 34 (Impaired clearance)"),
        ("AI Action: ", "ELEVATES TO CRITICAL ALERT. Identifies CYP2C9 metabolic blockage multiplying bleeding risk."),
        ("Recommendation: ", "Switch to Apixaban (DOAC) 5mg BID or reduce Warfarin by 50%.")
    ]
    for b_pre, txt in points_a:
        p2 = ca_tf.add_paragraph()
        p2.space_before = Pt(8)
        r1 = p2.add_run()
        r1.text = b_pre
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = C_NAVY
        r2 = p2.add_run()
        r2.text = txt
        r2.font.size = Pt(9)
        r2.font.color.rgb = C_SLATE

    # Card B: Benign Suppressed Alert
    cb = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, Inches(1.5), col_w, Inches(5.4))
    cb.fill.solid()
    cb.fill.fore_color.rgb = C_WHITE
    cb.line.color.rgb = RGBColor(187, 247, 208)
    cb_tf = cb.text_frame
    cb_tf.word_wrap = True
    cb_tf.margin_top = Inches(0.3)
    cb_tf.margin_left = Inches(0.25)
    cb_tf.margin_right = Inches(0.25)

    p = cb_tf.paragraphs[0]
    p.text = "🟢 CASE 2: SUPPRESSED ALERT"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    points_b = [
        ("Patient: ", "Marcus Chen (45M), Primary Hyperlipidemia"),
        ("Active Regimen: ", "Atorvastatin 20mg daily"),
        ("Prescribed: ", "Azithromycin 500mg (5-day course)"),
        ("Lab Context: ", "eGFR 95 (Normal), ALT 22 (Normal), AST 19 (Normal), CK 85"),
        ("AI Action: ", "ALERT FATIGUE SUPPRESSED. Flags that Azithromycin is an azalide with negligible CYP3A4 inhibition."),
        ("Clinical Result: ", "Clinician is spared an unnecessary warning; safe to proceed with standard prescription.")
    ]
    for b_pre, txt in points_b:
        p2 = cb_tf.add_paragraph()
        p2.space_before = Pt(8)
        r1 = p2.add_run()
        r1.text = b_pre
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = C_NAVY
        r2 = p2.add_run()
        r2.text = txt
        r2.font.size = Pt(9)
        r2.font.color.rgb = C_SLATE

    # Card C: Action Bar & Governance
    cc = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c3_left, Inches(1.5), col_w, Inches(5.4))
    cc.fill.solid()
    cc.fill.fore_color.rgb = C_WHITE
    cc.line.color.rgb = C_BORDER
    cc_tf = cc.text_frame
    cc_tf.word_wrap = True
    cc_tf.margin_top = Inches(0.3)
    cc_tf.margin_left = Inches(0.25)
    cc_tf.margin_right = Inches(0.25)

    p = cc_tf.paragraphs[0]
    p.text = "⚖️ ACTION BAR & GOVERNANCE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    points_c = [
        ("1-Click Accept Alternate: ", "Instant substitution with recommended safe therapy (e.g. Apixaban)."),
        ("Dosage Titration Tool: ", "Adjusts and documents dose modifications (e.g. 50% Warfarin cut)."),
        ("Structured Override: ", "Mandatory clinical justification capture when overriding critical warnings."),
        ("HIPAA Audit Trail: ", "Immutable timestamps, patient MRN, risk tier, action, and clinician ID."),
        ("Interoperable Export: ", "Live CSV & JSON export capabilities for EHR integration and clinical audits.")
    ]
    for b_pre, txt in points_c:
        p2 = cc_tf.add_paragraph()
        p2.space_before = Pt(8)
        r1 = p2.add_run()
        r1.text = b_pre
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = C_NAVY
        r2 = p2.add_run()
        r2.text = txt
        r2.font.size = Pt(9)
        r2.font.color.rgb = C_SLATE

    # Save presentation
    prs.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    create_deck()
